"""Guía de correspondencias: qué texto original se ha sustituido por cada alias.

Es un registro para el bufete. CONTIENE LOS DATOS REALES: se guarda solo en local, junto
al documento anonimizado, con un nombre y un aviso que dejan claro que no se comparte.
"""

from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

from anonimizador.core.lang import Language
from anonimizador.core.matcher import Match, Term
from anonimizador.i18n import tr

USER_SOURCE = "usuario"


@dataclass
class GuideEntry:
    replacement: str
    source: str  # "usuario" o tipo de dato detectado ("DNI", "PERSONA"...)
    originals: list[str] = field(default_factory=list)  # lo que se pidió sustituir
    found: Counter = field(default_factory=Counter)  # forma encontrada -> veces

    @property
    def count(self) -> int:
        return sum(self.found.values())


def build_entries(terms: list[Term], matches: list[Match]) -> list[GuideEntry]:
    """Una entrada por alias, en el orden de la lista (primero los del usuario)."""
    entries: dict[str, GuideEntry] = {}
    for term in terms:
        entry = entries.setdefault(term.replacement, GuideEntry(term.replacement, term.source))
        if term.original not in entry.originals:
            entry.originals.append(term.original)
    for match in matches:
        entry = entries.setdefault(match.replacement, GuideEntry(match.replacement, match.source))
        entry.found[match.text] += 1
    return sorted(entries.values(), key=lambda e: e.source != USER_SOURCE)


def default_path(output: Path, language: Language) -> Path:
    return output.with_name(output.stem + tr(language.code, "guide_suffix") + ".docx")


def write_guide(
    path: Path, entries: list[GuideEntry], source_doc: Path, output_doc: Path, language: Language
) -> Path:
    lang = language.code
    doc = Document()
    doc.core_properties.author = ""
    doc.add_heading(tr(lang, "guide_title"), level=1)

    warning = doc.add_paragraph().add_run(tr(lang, "guide_warning"))
    warning.bold = True
    warning.font.highlight_color = WD_COLOR_INDEX.YELLOW

    for key, value in (
        ("guide_source", source_doc.name),
        ("guide_output", output_doc.name),
        ("guide_date", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ):
        p = doc.add_paragraph()
        p.add_run(f"{tr(lang, key)}: ").bold = True
        p.add_run(value)

    headers = ["guide_col_alias", "guide_col_original", "guide_col_found", "guide_col_type",
               "guide_col_count"]  # fmt: skip
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, key in zip(table.rows[0].cells, headers, strict=True):
        cell.text = ""
        cell.paragraphs[0].add_run(tr(lang, key)).bold = True

    for entry in entries:
        row = table.add_row().cells
        row[0].text = entry.replacement
        row[1].text = "\n".join(entry.originals)
        row[2].text = (
            "\n".join(f"{form} ({n})" for form, n in entry.found.most_common())
            if entry.found
            else tr(lang, "guide_not_found")
        )
        row[3].text = (
            tr(lang, "guide_type_user")
            if entry.source == USER_SOURCE
            else tr(lang, f"label_{entry.source}")
        )
        row[4].text = str(entry.count)

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.save(str(path))
    return path
