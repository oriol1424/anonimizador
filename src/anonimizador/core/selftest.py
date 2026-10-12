"""Autoprueba del paquete: `Anonimitzador --prova [fitxer_resultat.txt]`.

Crea un Word y un PDF inventados (con una página escaneada) en una carpeta temporal, los
anonimiza y comprueba que no queda nada y que el OCR y el detector de nombres funcionan.
Sirve para verificar el paquete en GitHub Actions y en los equipos del bufete. No usa red
ni documentos reales; no escribe nada fuera de la carpeta temporal (salvo el resultado).
"""

import os
import platform
import subprocess
import tempfile
import time
from pathlib import Path

import pymupdf
from docx import Document

from anonimizador import __version__
from anonimizador.core import ner, pdf_io
from anonimizador.core.lang import CATALAN
from anonimizador.core.matcher import Term
from anonimizador.core.pipeline import analyze, anonymize

NAME = "Joan Puig i Ferrer"
SECRETS = ("Puig", "12345678Z", "F-123456-X", "+376 812 345")
LINES = [
    f"En {NAME}, amb DNI 12345678Z i NRT F-123456-X, compareix davant la Batllia.",
    "Telèfon +376 812 345. El perit Marcel·lí Roca ho confirma.",
]


def _word(path: Path) -> Path:
    doc = Document()
    for line in LINES:
        doc.add_paragraph(line)
    doc.save(str(path))
    return path


def _pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(LINES):
        page.insert_text((56, 72 + i * 18), line, fontsize=10)
    scan = pymupdf.open().new_page()
    for i, line in enumerate(LINES):
        scan.insert_text((56, 90 + i * 22), line, fontsize=12, fontname="tiro")
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842),
                                pixmap=scan.get_pixmap(dpi=150))  # fmt: skip
    doc.save(str(path))
    return path


def _readable(path: Path) -> str:
    if path.suffix == ".docx":
        return "\n".join(p.text for p in Document(str(path)).paragraphs)
    if path.suffix == ".pdf":
        return pdf_io.extract_text(path)  # con OCR: también la página escaneada
    return path.read_text(encoding="utf-8")


def _cpu() -> str:
    """Modelo de procesador (el detector int8 depende de sus instrucciones)."""
    try:
        if platform.system() == "Darwin":
            out = subprocess.run(["sysctl", "-n", "machdep.cpu.brand_string"],
                                 capture_output=True, text=True, timeout=5)  # fmt: skip
            return out.stdout.strip() or platform.processor()
        if platform.system() == "Windows":
            import winreg

            cpu_key = "HARDWARE/DESCRIPTION/System/CentralProcessor/0".replace("/", "\\")
            key = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, cpu_key)
            return str(winreg.QueryValueEx(key, "ProcessorNameString")[0]).strip()
    except (OSError, subprocess.SubprocessError):
        pass
    return platform.processor() or "?"


def _ner_diagnosis() -> str:
    """Qué ve el detector en la frase de prueba (texto inventado: se puede mostrar)."""
    import numpy as np

    model = ner._get_model()
    enc = model._encodings(LINES[1])[0]
    feed = {"input_ids": np.array([enc.ids], dtype=np.int64)}
    if "attention_mask" in model.inputs:
        feed["attention_mask"] = np.array([enc.attention_mask], dtype=np.int64)
    logits = model.session.run(None, feed)[0][0]
    tags = [model.id2label[int(i)] for i in logits.argmax(-1)]
    return f"etiquetes: {', '.join(sorted(set(tags)))}; màx {float(np.abs(logits).max()):.1f}"


def run(window_check=None) -> tuple[bool, list[str]]:
    lines = [
        f"Anonimitzador {__version__} · {platform.system()} {platform.release()} "
        f"({platform.machine()})",
        f"Processador: {_cpu()} · {os.cpu_count()} nuclis",
    ]
    ok = True
    if not ner.is_available():
        ok = False
        lines.append("ERROR: el detector de noms (model) no s'ha trobat")
    elif not ner.is_working():
        ok = False
        lines.append(f"ERROR: el detector de noms no funciona en aquest processador "
                     f"({_ner_diagnosis()})")  # fmt: skip
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        for src, outputs in ((_word(folder / "prova.docx"), (".docx", ".md")),
                             (_pdf(folder / "prova.pdf"), (".pdf", ".md"))):  # fmt: skip
            start = time.monotonic()
            try:
                terms = [Term(NAME, "[CLIENT_1]")]
                analysis = analyze(src, terms, CATALAN)
                for suggestion in analysis.suggestions:
                    terms.extend(suggestion.to_terms())
                for suffix in outputs:
                    out = folder / f"{src.stem}_sortida{suffix}"
                    report = anonymize(src, out, terms)
                    left = [s for s in SECRETS if s in _readable(out)]
                    leak = [n.kind for n in report.notices if n.is_leak]
                    status = "OK" if not left and not leak else "ERROR"
                    ok &= status == "OK"
                    detail = f" (queda: {', '.join(left + leak)})" if status != "OK" else ""
                    lines.append(f"{status}: {src.suffix} -> {suffix}{detail}")
                found = {s.label for s in analysis.suggestions}
                if ner.is_available() and "PERSONA" not in found:
                    ok = False
                    lines.append(
                        f"ERROR: {src.suffix}: el detector no ha trobat cap persona "
                        f"({_ner_diagnosis()})"
                    )
            except Exception as error:  # noqa: BLE001 - es informa i es continua
                ok = False
                lines.append(f"ERROR: {src.suffix}: {type(error).__name__}")
            lines.append(f"   ({time.monotonic() - start:.1f} s)")
    if window_check is not None:
        try:
            lines.append(window_check())
        except Exception as error:  # noqa: BLE001
            ok = False
            lines.append(f"ERROR: finestra: {type(error).__name__}")
    lines.append("RESULTAT: OK" if ok else "RESULTAT: ERROR")
    return ok, lines


def main(result_file: str | None = None, window_check=None) -> int:
    ok, lines = run(window_check)
    text = "\n".join(lines) + "\n"
    if result_file:
        Path(result_file).write_text(text, encoding="utf-8")
    try:
        print(text, end="")
    except (OSError, AttributeError):  # app sin consola (Windows): solo el fichero
        pass
    return 0 if ok else 1
