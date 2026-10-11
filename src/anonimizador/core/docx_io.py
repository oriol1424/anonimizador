"""Lectura y anonimización de documentos Word (.docx) conservando el formato.

Se trabaja directamente sobre el XML para cubrir todo lo que puede contener datos:
cuerpo, tablas, cuadros de texto, cabeceras, pies, notas, comentarios, control de
cambios (texto borrado y autores), códigos de campo, hipervínculos y metadatos.
"""

import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.part import Part, XmlPart
from lxml import etree

from anonimizador.core import images
from anonimizador.core.matcher import Match, Matcher
from anonimizador.core.progress import PageProgress, Steps
from anonimizador.core.result import ProcessResult

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
P, R = f"{{{W}}}p", f"{{{W}}}r"
T, DEL_T, INSTR_T = f"{{{W}}}t", f"{{{W}}}delText", f"{{{W}}}instrText"
TAB, BR, CR = f"{{{W}}}tab", f"{{{W}}}br", f"{{{W}}}cr"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A_P, A_T, A_BR = f"{{{A}}}p", f"{{{A}}}t", f"{{{A}}}br"
C_V = "{http://schemas.openxmlformats.org/drawingml/2006/chart}v"  # textos de gráficos
DOC_VAR = f"{{{W}}}docVar"

_WP = "application/vnd.openxmlformats-officedocument.wordprocessingml."
TEXT_CONTENT_TYPES = {
    _WP + "document.main+xml",
    _WP + "template.main+xml",
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    _WP + "header+xml",
    _WP + "footer+xml",
    _WP + "footnotes+xml",
    _WP + "endnotes+xml",
    _WP + "comments+xml",
    _WP + "document.glossary+xml",  # bloques de creación (textos reutilizables)
}
# Gráficos y SmartArt: texto DrawingML (a:p/a:t) y valores de los gráficos (c:v).
DRAWING_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
    "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
    "application/vnd.ms-office.drawingml.diagramDrawing+xml",
}
SETTINGS_CONTENT_TYPE = _WP + "settings+xml"  # variables de documento (docVars)
CUSTOM_XML = re.compile(r"^/customXml/item\d+\.xml$")  # datos de controles de contenido
PEOPLE_CONTENT_TYPE = _WP + "people+xml"
CUSTOM_PROPS = "/docProps/custom.xml"  # propiedades personalizadas ("Client = ...")
THUMBNAIL_REL = "/metadata/thumbnail"  # miniatura de la 1ª página (imagen con nombres)
# Texto alternativo de imágenes y formas: `descr`/`title` (DrawingML) y `alt` (VML).
ALT_TEXT_ATTRS = {"descr", "title", "alt"}
CORE_TEXT_FIELDS = ("title", "subject", "keywords", "comments", "category", "content_status")


@dataclass
class _Segment:
    node: etree._Element | None  # None = separador virtual (tabulador, salto de línea)
    start: int
    end: int


class _XmlRoot:
    """Raíz XML de una parte del paquete; guarda los cambios al cerrar."""

    def __init__(self, part: Part):
        self.part = part
        self.is_xml_part = isinstance(part, XmlPart)
        self.root = part._element if self.is_xml_part else etree.fromstring(part.blob)

    def commit(self) -> None:
        if not self.is_xml_part:
            self.part._blob = etree.tostring(
                self.root, xml_declaration=True, encoding="UTF-8", standalone=True
            )


def _xml_roots(doc: DocxDocument, content_types: set[str]) -> Iterator[_XmlRoot]:
    for part in doc.part.package.iter_parts():
        if part.content_type in content_types:
            yield _XmlRoot(part)


def _owner_paragraph(el: etree._Element) -> etree._Element | None:
    parent = el.getparent()
    while parent is not None and parent.tag != P:
        parent = parent.getparent()
    return parent


def _paragraph_segments(p: etree._Element) -> tuple[str, list[_Segment]]:
    """Texto de un párrafo y de qué nodo XML sale cada tramo.

    Un nombre puede estar partido en varios runs ("Pé" + "rez" con formatos distintos);
    al concatenar los nodos lo encontramos igual.
    """
    pieces: list[str] = []
    segments: list[_Segment] = []
    pos = 0
    for el in p.iter(T, DEL_T, TAB, BR, CR):
        if _owner_paragraph(el) is not p:
            continue  # pertenece a un cuadro de texto anidado; se procesa aparte
        if el.tag in (T, DEL_T):
            text, node = el.text or "", el
        elif el.getparent().tag == R:
            text, node = " ", None
        else:
            continue  # definición de tabulaciones en las propiedades del párrafo
        segments.append(_Segment(node, pos, pos + len(text)))
        pieces.append(text)
        pos += len(text)
    return "".join(pieces), segments


def _drawing_segments(p: etree._Element) -> tuple[str, list[_Segment]]:
    """Como `_paragraph_segments`, para un párrafo DrawingML (gráficos, SmartArt)."""
    pieces: list[str] = []
    segments: list[_Segment] = []
    pos = 0
    for el in p.iter(A_T, A_BR):
        text, node = (el.text or "", el) if el.tag == A_T else (" ", None)
        segments.append(_Segment(node, pos, pos + len(text)))
        pieces.append(text)
        pos += len(text)
    return "".join(pieces), segments


def _field_segments(p: etree._Element) -> tuple[str, list[_Segment]]:
    """Códigos de campo de un párrafo, unidos: Word los parte a menudo en varios trozos
    (`HYPERLINK "mailto:joan` + `.puig@..."`)."""
    pieces: list[str] = []
    segments: list[_Segment] = []
    pos = 0
    for el in p.iter(INSTR_T):
        if _owner_paragraph(el) is not p:
            continue
        text = el.text or ""
        segments.append(_Segment(el, pos, pos + len(text)))
        pieces.append(text)
        pos += len(text)
    return "".join(pieces), segments


def _segmented(root: etree._Element) -> Iterator[tuple[str, list[_Segment]]]:
    """Todos los textos de una parte que se buscan como un párrafo: párrafos de Word,
    sus códigos de campo y párrafos DrawingML."""
    for p in root.iter(P):
        yield _paragraph_segments(p)
        yield _field_segments(p)
    for p in root.iter(A_P):
        yield _drawing_segments(p)


def _loose_texts(doc: DocxDocument) -> Iterator[tuple[etree._Element, str | None, _XmlRoot]]:
    """Textos sueltos fuera de los párrafos: (elemento, atributo o None = su texto, parte).

    Valores de gráficos, datos de controles de contenido (customXml) y variables de
    documento."""
    for part in doc.part.package.iter_parts():
        is_custom_xml = bool(CUSTOM_XML.match(str(part.partname)))
        if not (is_custom_xml or part.content_type in DRAWING_CONTENT_TYPES
                or part.content_type == SETTINGS_CONTENT_TYPE):  # fmt: skip
            continue
        handle = _XmlRoot(part)
        for el in handle.root.iter(tag=etree.Element):
            if el.tag == DOC_VAR:
                for attr in el.attrib:
                    if etree.QName(attr).localname == "val":
                        yield el, attr, handle
            elif (el.tag == C_V or is_custom_xml) and el.text and el.text.strip():
                yield el, None, handle


def _apply_matches(segments: list[_Segment], matches: list[Match]) -> None:
    # De atrás hacia delante para que los offsets anteriores sigan siendo válidos.
    for m in reversed(matches):
        first = True
        for seg in segments:
            if seg.node is None or seg.end <= m.start or seg.start >= m.end:
                continue
            current = seg.node.text or ""
            if first:
                tail = current[m.end - seg.start :] if m.end <= seg.end else ""
                new = current[: m.start - seg.start] + m.replacement + tail
                first = False
            else:
                new = current[m.end - seg.start :] if m.end < seg.end else ""
            seg.node.text = new
            seg.node.set(XML_SPACE, "preserve")


def _anonymize_authors(root: etree._Element) -> None:
    """Autores del control de cambios y comentarios (suelen ser abogados del bufete)."""
    for el in root.iter(tag=etree.Element):
        for attr in list(el.attrib):
            name = etree.QName(attr).localname
            if name == "author":
                el.set(attr, "Autor")
            elif name == "initials":
                el.set(attr, "A")
            elif name in ("userId", "providerId"):
                el.set(attr, "")


def _alt_texts(root: etree._Element) -> Iterator[tuple[etree._Element, str]]:
    for el in root.iter(tag=etree.Element):
        for attr in el.attrib:
            if etree.QName(attr).localname in ALT_TEXT_ATTRS and el.get(attr):
                yield el, attr


def _custom_props(doc: DocxDocument) -> _XmlRoot | None:
    for part in doc.part.package.iter_parts():
        if str(part.partname) == CUSTOM_PROPS:
            return _XmlRoot(part)
    return None


def _remove_thumbnail(doc: DocxDocument) -> None:
    """La miniatura es una imagen de la portada: se quita del paquete."""
    rels = doc.part.package.rels
    for r_id in [k for k, rel in rels.items() if rel.reltype.endswith(THUMBNAIL_REL)]:
        rels.pop(r_id)


def _anonymize_metadata(doc: DocxDocument, matcher: Matcher) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    for field in CORE_TEXT_FIELDS:
        value = getattr(props, field)
        if value:
            setattr(props, field, matcher.replace(value))
    for part in doc.part.package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            handle = _XmlRoot(part)
            for el in handle.root.iter(tag=etree.Element):
                if etree.QName(el).localname in ("Company", "Manager", "HyperlinkBase"):
                    el.text = ""
            handle.commit()
    custom = _custom_props(doc)
    if custom is not None:
        for el in custom.root.iter(tag=etree.Element):
            if el.text and el.text.strip():
                el.text = matcher.replace(el.text)
        custom.commit()
    _remove_thumbnail(doc)


def _anonymize_links(doc: DocxDocument, matcher: Matcher) -> None:
    for part in doc.part.package.iter_parts():
        for rel in part.rels.values():
            if rel.is_external:
                new_target = matcher.replace(rel.target_ref)
                if new_target != rel.target_ref:
                    rel._target = new_target


def _media(doc: DocxDocument) -> list[Part]:
    return [p for p in doc.part.package.iter_parts() if str(p.partname).startswith("/word/media/")]


def _anonymize_images(doc: DocxDocument, matcher: Matcher, progress: PageProgress | None):
    """Imágenes del documento (fotos de DNI, escaneados pegados): OCR y borrado.

    Devuelve (sustituciones, datos que aún se leen, imágenes leídas, imágenes ilegibles)."""
    matches: list[Match] = []
    leftovers: list[Match] = []
    read = unreadable = 0
    media = _media(doc)
    steps = Steps(progress, len(media))
    for part in media:
        pix = images.load(part.blob)
        if pix is None:
            unreadable += 1  # EMF, WMF, SVG...: no se puede leer
        elif images.worth_reading(pix):
            read += 1
            result = images.anonymize(part.blob, part.partname.ext, matcher)
            if result.data is not None:
                part._blob = result.data
            matches += result.matches
            leftovers += result.leftovers
        steps.tick()
    return matches, leftovers, read, unreadable


def anonymize(
    src: Path, dst: Path, matcher: Matcher, progress: PageProgress | None = None
) -> ProcessResult:
    """`progress` cuenta las imágenes leídas con OCR (lo único lento)."""
    doc = Document(str(src))
    all_matches: list[Match] = []

    for handle in _xml_roots(doc, TEXT_CONTENT_TYPES | DRAWING_CONTENT_TYPES):
        for text, segments in _segmented(handle.root):
            matches = matcher.find(text)
            if matches:
                _apply_matches(segments, matches)
                all_matches.extend(matches)
        for el, attr in _alt_texts(handle.root):
            el.set(attr, matcher.replace(el.get(attr)))
        _anonymize_authors(handle.root)
        handle.commit()

    for handle in _xml_roots(doc, {PEOPLE_CONTENT_TYPE}):
        _anonymize_authors(handle.root)
        handle.commit()

    changed: dict[int, _XmlRoot] = {}
    for el, attr, handle in _loose_texts(doc):
        value = el.get(attr) if attr else el.text
        new = matcher.replace(value)
        if new != value:
            el.set(attr, new) if attr else setattr(el, "text", new)
            changed[id(handle.part)] = handle
    for handle in changed.values():
        handle.commit()

    image_matches, image_leftovers, read, unreadable = _anonymize_images(doc, matcher, progress)
    all_matches.extend(image_matches)

    _anonymize_metadata(doc, matcher)
    _anonymize_links(doc, matcher)
    doc.save(str(dst))
    return ProcessResult(
        all_matches, unreadable + count_embedded(doc), images_ocr=read,
        image_leftovers=image_leftovers,
    )  # fmt: skip


def count_embedded(doc: DocxDocument) -> int:
    """Objetos incrustados (hojas de cálculo, otros documentos): no se revisan."""
    return sum(
        1 for part in doc.part.package.iter_parts()
        if str(part.partname).startswith("/word/embeddings/")
    )  # fmt: skip


def extract_text(path: Path, progress: PageProgress | None = None) -> str:
    """Todo el texto susceptible de contener datos, un párrafo por línea (también el de
    las imágenes, leído con OCR; `progress` cuenta las imágenes)."""
    doc = Document(str(path))
    lines: list[str] = []
    media = _media(doc)
    steps = Steps(progress, len(media))
    for part in media:
        lines.append(images.text_of(part.blob))
        steps.tick()
    for handle in _xml_roots(doc, TEXT_CONTENT_TYPES | DRAWING_CONTENT_TYPES):
        lines.extend(text for text, _ in _segmented(handle.root) if text.strip())
        lines.extend(el.get(attr) for el, attr in _alt_texts(handle.root))
    custom = _custom_props(doc)
    if custom is not None:
        lines.extend(el.text for el in custom.root.iter(tag=etree.Element)
                     if el.text and el.text.strip())  # fmt: skip
    lines.extend(el.get(attr) if attr else el.text for el, attr, _ in _loose_texts(doc))
    props = doc.core_properties
    lines.extend(
        v for f in ("author", "last_modified_by", *CORE_TEXT_FIELDS) if (v := getattr(props, f))
    )
    for part in doc.part.package.iter_parts():
        lines.extend(r.target_ref for r in part.rels.values() if r.is_external)
    return "\n".join(lines)


def visible_text(path: Path) -> str:
    """Texto de todas las partes con párrafos (cuerpo, cabeceras, notas, comentarios,
    gráficos), sin OCR de imágenes: para la revisión final."""
    doc = Document(str(path))
    lines: list[str] = []
    for handle in _xml_roots(doc, TEXT_CONTENT_TYPES | DRAWING_CONTENT_TYPES):
        lines.extend(text for text, _ in _segmented(handle.root) if text.strip())
    return "\n".join(lines)


def stats(path: Path) -> tuple[int, int]:
    """(caracteres del texto, imágenes que se leerán con OCR): para estimar el tiempo."""
    doc = Document(str(path))
    chars = sum(
        len(text)
        for handle in _xml_roots(doc, TEXT_CONTENT_TYPES | DRAWING_CONTENT_TYPES)
        for text, _ in _segmented(handle.root)
    )  # fmt: skip
    readable = 0
    for part in _media(doc):
        pix = images.load(part.blob)
        readable += pix is not None and images.worth_reading(pix)
    return chars, readable


def _heading_level(p) -> int:
    """Nivel de título del párrafo (0 = no es título), por su estilo de Word."""
    name = (p.style.name if p.style is not None else "").lower()
    if name in ("title", "títol", "título"):
        return 1
    for prefix in ("heading ", "títol ", "título "):
        if name.startswith(prefix) and name[len(prefix) :].isdigit():
            return min(6, int(name[len(prefix) :]))
    return 0


def _is_list(p) -> bool:
    name = (p.style.name if p.style is not None else "").lower()
    return p._p.pPr is not None and p._p.pPr.numPr is not None or name.startswith("list")


def _md_inline(p) -> str:
    """Texto del párrafo con negrita (**) y cursiva (*) de Word."""
    pieces: list[str] = []
    runs = []
    for item in p.iter_inner_content():  # los enlaces también tienen texto
        runs.extend(item.runs if hasattr(item, "runs") else [item])
    for run in runs:
        text = run.text.replace("*", "\\*")
        if not text.strip():
            pieces.append(text)
            continue
        lead = text[: len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()) :]
        core = text.strip()
        if run.bold and run.italic:
            core = f"***{core}***"
        elif run.bold:
            core = f"**{core}**"
        elif run.italic:
            core = f"*{core}*"
        pieces.append(lead + core + trail)
    return "".join(pieces).replace("****", "").strip()


def markdown(path: Path) -> str:
    """El documento (ya anonimizado) en Markdown: títulos, listas, negritas y tablas. Es
    el formato que mejor entiende una IA."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == P:
            p = Paragraph(child, doc)
            text = _md_inline(p)
            if not text:
                continue
            level = _heading_level(p)
            if level:
                blocks.append("#" * level + " " + p.text.strip())
            elif _is_list(p):
                blocks.append("- " + text)
            else:
                blocks.append(text)
        elif child.tag == f"{{{W}}}tbl":
            rows = [[c.text.strip().replace("|", "\\|").replace("\n", " ") for c in r.cells]
                    for r in Table(child, doc).rows]  # fmt: skip
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            table = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
            table += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            blocks.append("\n".join(table))
    # Las listas seguidas van juntas; el resto, separado por una línea en blanco.
    out: list[str] = []
    for block in blocks:
        if out and block.startswith("- ") and out[-1].startswith("- "):
            out[-1] += "\n" + block
        else:
            out.append(block)
    return "\n\n".join(out).strip() + "\n"


def body_text(path: Path) -> str:
    """Texto visible del cuerpo del documento (para exportar a .txt).

    Excluye el texto borrado del control de cambios, comentarios, cabeceras y pies.
    """
    doc = Document(str(path))
    lines = []
    for p in doc.part.element.iter(P):
        lines.append(
            "".join(
                (el.text or "") if el.tag == T else "\t"
                for el in p.iter(T, TAB)
                if _owner_paragraph(el) is p and (el.tag == T or el.getparent().tag == R)
            )
        )
    return "\n".join(lines).strip() + "\n"
