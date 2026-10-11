"""Pruebas de la ventana (sin pantalla: QT_QPA_PLATFORM=offscreen en conftest)."""

import json
import threading
import time
from pathlib import Path

import pytest
from docx import Document
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QMessageBox

from anonimizador.core.lang import CATALAN, fill_preset
from anonimizador.core.matcher import Term
from anonimizador.core.pipeline import Analysis, Notice, Report, Suggestion, analyze
from anonimizador.core.progress import Cancelled
from anonimizador.ui import main_window as main_window_module
from anonimizador.ui.main_window import (
    STEP_DOCUMENT,
    STEP_SAVE,
    STEP_SUGGESTIONS,
    STEP_TERMS,
    MainWindow,
)
from anonimizador.ui.worker import run_in_background


@pytest.fixture
def window(qtbot, tmp_path):
    w = MainWindow(
        QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat),
        keep_path=tmp_path / "no_anonimitzar.json",  # nunca la configuración real
    )
    qtbot.addWidget(w)
    return w


def _fill_list(window, rows):
    window.table.clear_rows()
    for original, replacement in rows:
        window.table.add_row(original, replacement)


def test_fill_preset():
    assert fill_preset("[CLIENT_{n}]", set()) == "[CLIENT_1]"
    assert fill_preset("[CLIENT_{n}]", {"[CLIENT_1]", "[CLIENT_2]"}) == "[CLIENT_3]"
    assert fill_preset("***", {"***"}) == "***"


def test_starts_in_catalan_on_the_first_step(window):
    assert window.windowTitle() == "Anonimitzador"
    assert window.language.code == "ca"
    assert window.current_step() == STEP_DOCUMENT
    assert window.run_button.text() == "Anonimitza"
    assert window.table.combo(0).value() == "[PERSONA_1]"
    assert window.sidebar.steps[0].text_label.text() == "Document"


def test_switch_to_spanish(window):
    window.sidebar.language.buttons["es"].click()
    assert window.windowTitle() == "Anonimizador"
    assert window.run_button.text() == "Anonimizar"
    assert window.sidebar.steps[1].text_label.text() == "Tu lista"
    assert window.table.combo(0).itemText(1) == "[CLIENTE_n]"
    assert window.table.combo(0).value() == "[PERSONA_1]"  # no se pierde lo escrito


def test_presets_get_next_free_number(window):
    window.table.add_row()
    first, second = window.table.combo(0), window.table.combo(1)
    first.choose("[CLIENT_{n}]")
    second.choose("[CLIENT_{n}]")
    assert (first.value(), second.value()) == ("[CLIENT_1]", "[CLIENT_2]")


def test_steps_unlock_in_order(window, demanda_ca: Path):
    window.go_to(STEP_TERMS)
    assert window.current_step() == STEP_DOCUMENT  # sin documento no se puede avanzar
    window.set_input(demanda_ca)
    assert window.document_page.file_name.text() == "demanda_ca.docx"
    assert window.document_page.continue_button.isEnabled()
    window.go_to(STEP_TERMS)
    assert window.current_step() == STEP_TERMS
    window.go_to(STEP_SUGGESTIONS)
    assert window.current_step() == STEP_TERMS  # falta analizar


def test_changing_the_list_requires_a_new_analysis(window, demanda_ca: Path):
    window.set_input(demanda_ca)
    _fill_list(window, [("Joan Puig i Ferrer", "[CLIENT_1]")])
    terms = [Term("Joan Puig i Ferrer", "[CLIENT_1]")]
    window._analysis_done(analyze(demanda_ca, terms, CATALAN), window._current_key(), terms)
    assert window.current_step() == STEP_SUGGESTIONS
    assert window._reachable()[STEP_SUGGESTIONS]
    window.table.add_row("Anna Vila", "[PERSONA_2]")
    assert not window._reachable()[STEP_SUGGESTIONS]


def test_input_sets_default_output(window, demanda_ca: Path):
    window.set_input(demanda_ca)
    assert window.folder_edit.text() == str(demanda_ca.parent)
    assert window.name_edit.text() == "demanda_ca_anonimitzat"
    window.save_page.formats.set_value(".txt")
    assert window.output_path() == demanda_ca.parent / "demanda_ca_anonimitzat.txt"


def test_output_suffix_follows_language_if_untouched(window, demanda_ca: Path):
    window.set_input(demanda_ca)
    window.set_language("es")
    assert window.name_edit.text() == "demanda_ca_anonimizado"
    window.name_edit.setText("el_meu_nom")
    window.set_language("ca")
    assert window.name_edit.text() == "el_meu_nom"


def test_pdf_input_offers_pdf_and_txt(window, demanda_pdf: Path, demanda_ca: Path):
    window.set_input(demanda_pdf)
    assert list(window.save_page.formats.buttons) == [".pdf", ".txt"]
    assert "escanejades" in window.document_page.file_details.text()
    assert window.output_path().suffix == ".pdf"
    window.save_page.formats.set_value(".txt")
    window.set_input(demanda_ca)  # se conserva .txt al cambiar a un Word
    assert window.output_path().suffix == ".txt"


def test_encrypted_pdf_is_rejected_when_chosen(window, tmp_path, monkeypatch):
    import pymupdf

    path = tmp_path / "protegit.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "Secret")
    doc.save(str(path), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    errors = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: errors.append(a[2]))
    window.set_input(path)
    assert window.input_path is None
    assert errors and "contrasenya" in errors[0]


def test_save_and_load_list(window, tmp_path):
    _fill_list(window, [("Joan Puig i Ferrer", "[CLIENT_1]")])
    path = tmp_path / "cas.json"
    data = {"idioma": "ca", "termes": [{"original": o, "substitut": r}
                                       for o, r in window.table.entries()]}  # fmt: skip
    path.write_text(json.dumps(data), encoding="utf-8")
    window.table.clear_rows()
    window.load_list_file(path)
    assert window.table.entries() == [("Joan Puig i Ferrer", "[CLIENT_1]")]


# --- Tareas en segundo plano --------------------------------------------------------------


def test_background_callback_runs_in_ui_thread(qtbot):
    seen = {}
    main = threading.get_ident()

    def finished(value):
        seen["value"], seen["thread"] = value, threading.get_ident()

    steps = []

    def work(progress):
        progress("read", 1, 2)
        return threading.get_ident()

    run_in_background(work, finished, lambda e: None, lambda *a: steps.append(a))
    qtbot.waitUntil(lambda: "value" in seen, timeout=5000)
    assert steps == [("read", 1, 2)]
    assert seen["value"] != main  # la tarea corrió en otro hilo
    assert seen["thread"] == main  # el resultado llegó al hilo de la interfaz


def test_background_task_can_be_cancelled(qtbot):
    events = []

    def work(progress):
        for i in range(500):
            progress("detect", i, 500)
            time.sleep(0.01)
        return "acabat"

    task = run_in_background(work, events.append, events.append, None,
                             lambda: events.append("cancel·lat"))  # fmt: skip
    time.sleep(0.05)
    task.cancel()
    qtbot.waitUntil(lambda: bool(events), timeout=5000)
    assert events == ["cancel·lat"]


# --- Sugerencias ---------------------------------------------------------------------------


def _analysis() -> Analysis:
    return Analysis(
        [
            Suggestion("PERSONA", "[PERSONA_1]", ("Núria Serra",), 1, expand=True),
            Suggestion("DNI", "[DNI_1]", ("12345678Z",), 2),
            Suggestion("TELEFONO", "[TELEFON_1]", ("+376 812 345",), 1),
            Suggestion("DIRECCION", "[ADRECA_1]", ("Avinguda Meritxell, 25",), 1),
        ],
        [Notice("ocr_doubtful", count=2, pages="3"), Notice("unreadable_pages", pages="5")],
    )


def test_suggestions_search_and_filters(window):
    page = window.suggestions_page
    page.show_analysis(_analysis(), CATALAN, set())
    visible = lambda: [r for r in range(page.table.rowCount()) if not page.table.isRowHidden(r)]  # noqa: E731
    assert visible() == [0, 1, 2, 3]
    page.search.setText("nuria")  # sin acento ni mayúscula
    assert visible() == [0]
    page.search.setText("")
    page._choose_group("group_ids")
    assert visible() == [1]
    page._choose_group(None)
    assert visible() == [0, 1, 2, 3]
    assert page.count_label.text() == "4 de 4 marcades"


def test_suggestions_serious_notices_first(window):
    page = window.suggestions_page
    page.show_analysis(_analysis(), CATALAN, set())
    first = page.notices._layout.itemAt(0).widget()
    assert first.property("level") == "serious"


def test_select_none_only_affects_visible_rows(window):
    page = window.suggestions_page
    page.show_analysis(_analysis(), CATALAN, set())
    page.search.setText("12345678Z")
    page.none_button.click()
    page.search.setText("")
    sources = {t.source for t in page.selected_terms()}
    assert sources == {"PERSONA", "TELEFONO", "DIRECCION"}


def test_suggestions_return_checked_terms_with_edited_alias(window, demanda_ca: Path):
    analysis = analyze(demanda_ca, [], CATALAN)
    page = window.suggestions_page
    page.show_analysis(analysis, CATALAN, set())
    dni_row = next(i for i, s in enumerate(analysis.suggestions) if s.label == "DNI")
    page.combos[dni_row].setEditText("[DNI_CLIENT]")
    phone_row = next(i for i, s in enumerate(analysis.suggestions) if s.label == "TELEFONO")
    page.set_checked(phone_row, False)
    terms = page.selected_terms()
    assert Term("12345678Z", "[DNI_CLIENT]", expand=False, source="DNI") in terms
    assert not [t for t in terms if t.source == "TELEFONO"]


# --- Flujo completo -----------------------------------------------------------------------


def test_full_flow_from_window(window, qtbot, demanda_ca: Path):
    window.set_input(demanda_ca)
    window.go_to(STEP_TERMS)
    _fill_list(window, [("Joan Puig i Ferrer", "[CLIENT_1]")])
    window.terms_page.analyze_button.click()
    qtbot.waitUntil(lambda: window.current_step() == STEP_SUGGESTIONS, timeout=20000)
    assert window.sidebar.steps[2].badge.text() == str(len(window.analysis.suggestions))
    window.suggestions_page.apply_button.click()
    assert window.current_step() == STEP_SAVE
    window.run_button.click()
    qtbot.waitUntil(lambda: window.pages.currentWidget() is window.result_page, timeout=20000)

    out = demanda_ca.parent / "demanda_ca_anonimitzat.docx"
    assert window.result_page.report.output == out
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Puig" not in text and "12345678Z" not in text
    assert "[CLIENT_1]" in text and "[DNI_1]" in text
    assert window.result_page.headline.text()

    window.result_page.new_button.click()
    assert window.current_step() == STEP_DOCUMENT and window.input_path is None
    assert window.table.entries() == [("Joan Puig i Ferrer", "[CLIENT_1]")]  # se conserva


def test_cancel_analysis_goes_back_to_the_list(window, qtbot, demanda_ca: Path, monkeypatch):
    def slow_analyze(src, terms, language, progress):
        for i in range(1000):
            progress("detect", i, 1000)
            time.sleep(0.01)
        raise AssertionError("no s'havia de completar")

    monkeypatch.setattr(main_window_module, "analyze", slow_analyze)
    window.set_input(demanda_ca)
    window.go_to(STEP_TERMS)
    window.start_analysis()
    assert window.pages.currentWidget() is window.progress_page
    qtbot.waitUntil(lambda: window.progress_page.bar.value() > 0, timeout=5000)
    window.progress_page.cancel_button.click()
    assert window.progress_page.cancel_button.text() == "Cancel·lant…"
    qtbot.waitUntil(lambda: window.current_step() == STEP_TERMS, timeout=5000)
    assert window._task is None and window.analysis is None


def test_error_returns_to_step_and_shows_message(window, qtbot, demanda_ca, monkeypatch):
    def broken(*args):
        raise RuntimeError("disc ple")

    errors = []
    monkeypatch.setattr(main_window_module, "analyze", broken)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: errors.append(a[2]))
    window.set_input(demanda_ca)
    window.go_to(STEP_TERMS)
    window.start_analysis()
    qtbot.waitUntil(lambda: bool(errors), timeout=5000)
    assert window.current_step() == STEP_TERMS
    assert "disc ple" in errors[0]


def test_progress_page_shows_phases_and_pages(window):
    page = window.progress_page
    page.start("progress_anonymizing", ["anonymize", "verify"], "ca")
    page.update_progress("anonymize", 52, 80)
    assert page.position.text() == "Pàgina 52 de 80"
    assert page.percent.text() == "65%"
    assert page._rows[0].detail.text() == "pàgina 52 de 80"
    assert page._rows[1].detail.text() == "pendent"
    page.update_progress("verify", 1, 80)
    assert page._rows[0].detail.text() == "fet"


def test_result_page_highlights_serious_notices(window, tmp_path):
    notices = [Notice("images", count=1), Notice("leftovers", detail="Puig", count=1)]
    report = Report(tmp_path / "out.docx", {"[CLIENT_1]": 3}, notices)
    window.show_result(report)
    assert window.result_page.headline.text() == (
        "Atenció: revisa el document abans de fer-lo servir"
    )
    first = window.result_page.notices._layout.itemAt(0).widget()
    assert first.property("level") == "serious"


def test_cancelled_is_not_an_exception_subclass():
    """Si lo fuera, la protección de páginas que fallan (except Exception) lo atraparía."""
    assert not issubclass(Cancelled, Exception)


# --- Lista «No anonimitzar» ----------------------------------------------------------------


def test_keep_list_is_saved_and_shown(window, qtbot, tmp_path):
    assert window.sidebar.keep_button.text() == "No anonimitzar · 0"
    window.add_to_keep_list(["Montserrat Soler i Vidal", "Montserrat Soler i Vidal"])
    assert window.sidebar.keep_button.text() == "No anonimitzar · 1"
    other = MainWindow(
        QSettings(str(tmp_path / "s2.ini"), QSettings.Format.IniFormat),
        keep_path=tmp_path / "no_anonimitzar.json",
    )
    qtbot.addWidget(other)
    assert other.keep == ["Montserrat Soler i Vidal"]  # se recuerda al volver a abrir


def test_keep_dialog_items(qtbot):
    from anonimizador.ui.keep_dialog import KeepDialog

    dialog = KeepDialog(["Montserrat Soler i Vidal"], "ca")
    qtbot.addWidget(dialog)
    dialog.add_row("  Despatx Soler & Associats ")
    dialog.add_row("")
    assert dialog.items() == ["Montserrat Soler i Vidal", "Despatx Soler & Associats"]
    dialog.table.selectRow(0)
    dialog.remove_selected()
    assert dialog.items() == ["Despatx Soler & Associats"]


def test_never_anonymize_from_suggestions(window):
    analysis = Analysis([
        Suggestion("PERSONA", "[PERSONA_1]", ("Sra. Soler", "Montserrat Soler i Vidal"), 3,
                   expand=True),
        Suggestion("DNI", "[DNI_1]", ("12345678Z",), 1),
    ])  # fmt: skip
    page = window.suggestions_page
    page.show_analysis(analysis, CATALAN, set())
    assert not page.keep_button.isEnabled()
    page.table.selectRow(0)
    assert page.keep_button.isEnabled()
    page.keep_button.click()
    assert window.keep == ["Montserrat Soler i Vidal"]  # la forma más completa
    assert page.table.isRowHidden(0)
    assert [t.source for t in page.selected_terms()] == ["DNI"]
    assert page.count_label.text() == "1 de 1 marcades"


def test_keep_list_reaches_analysis_and_output(window, qtbot, demanda_ca: Path):
    window.add_to_keep_list(["Montserrat Soler i Vidal"])
    window.set_input(demanda_ca)
    window.go_to(STEP_TERMS)
    _fill_list(window, [("Joan Puig i Ferrer", "[CLIENT_1]")])
    window.terms_page.analyze_button.click()
    qtbot.waitUntil(lambda: window.current_step() == STEP_SUGGESTIONS, timeout=20000)
    forms = {f for s in window.analysis.suggestions for f in s.forms}
    assert "Montserrat Soler i Vidal" not in forms
    window.suggestions_page.apply_button.click()
    window.run_button.click()
    qtbot.waitUntil(lambda: window.pages.currentWidget() is window.result_page, timeout=20000)
    out = window.result_page.report.output
    assert "Montserrat Soler i Vidal" in Document(str(out)).sections[0].header.paragraphs[0].text


def test_editing_keep_list_requires_new_analysis(window, demanda_ca: Path):
    window.set_input(demanda_ca)
    terms = window._user_terms()
    window._analysis_done(analyze(demanda_ca, terms, CATALAN), window._current_key(), terms)
    assert window._reachable()[STEP_SUGGESTIONS]
    window._save_keep(["Anna Vila"])  # como si se hubiera editado en el diálogo
    assert not window._reachable()[STEP_SUGGESTIONS]
