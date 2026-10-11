"""Genera documentos de prueba con datos INVENTADOS.

Uso manual:  python tests/corpus/generar_corpus.py
    -> crea en tests/corpus/ la demanda en catalán (.docx y .pdf), en castellano (.docx) y
       una demanda andorrana ante la Batllia (.docx y .pdf)
Cubre los casos difíciles: nombres partidos en varios runs con distinto formato, tablas,
cabecera, pie, comentarios, control de cambios, hipervínculos mailto y metadatos.
"""

from pathlib import Path

import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_tracked_deletion(paragraph, text: str, author: str) -> None:
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "1")
    deletion.set(qn("w:author"), author)
    deletion.set(qn("w:date"), "2026-01-01T00:00:00Z")
    run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = text
    run.append(del_text)
    deletion.append(run)
    paragraph._p.append(deletion)


def build_demanda_ca(path: Path) -> Path:
    """Demanda en catalán: el caso habitual del bufete."""
    doc = Document()
    doc.core_properties.author = "Montserrat Soler i Vidal"
    doc.core_properties.last_modified_by = "Montserrat Soler i Vidal"
    doc.core_properties.title = "Demanda Joan Puig i Ferrer contra Construccions del Vallès SL"

    section = doc.sections[0]
    section.header.paragraphs[
        0
    ].text = "Despatx Soler & Associats — Lletrada Montserrat Soler i Vidal"
    section.footer.paragraphs[0].text = "Tel. 972 123 456 · info@soler-associats.example"

    doc.add_heading("AL JUTJAT DE PRIMERA INSTÀNCIA", level=1)

    p = doc.add_paragraph("En ")
    p.add_run("Jo").bold = True
    p.add_run("an Puig").bold = True
    p.add_run(" i Ferrer").italic = True
    p.add_run(", major d’edat, amb DNI 12345678Z i domicili al Carrer Major 5, Girona, ")
    p.add_run("actuant en nom propi, davant el Jutjat compareix i DIU:")

    doc.add_paragraph(
        "Que el Sr. Puig va subscriure un contracte amb CONSTRUCCIONS DEL VALLÈS SL "
        "(NIF B12345674) en el procediment 1234/2025. J. Puig Ferrer va abonar les quantitats "
        "al compte ES91 2100 0418 4502 0005 1332."
    )
    doc.add_paragraph(
        "La part contrària, Martí d’Orriols, administrador de la societat, va reconèixer "
        "el deute. L’Orriols no ha pagat. El perit Marcel·lí Roca ho confirma; "
        "el senyor Roca ratificarà l’informe."
    )

    p = doc.add_paragraph("Contacte del demandant:	")
    p.add_run("Joan")
    p.add_run().add_tab()
    p.add_run("Puig i Ferrer")

    p = doc.add_paragraph("Correu: ")
    _add_hyperlink(p, "mailto:joan.puig@correu.example", "joan.puig@correu.example")

    p = doc.add_paragraph("Text revisat. ")
    _add_tracked_deletion(p, "Abans hi deia Joan Puig i Ferrer", "Montserrat Soler i Vidal")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Demandat"
    table.cell(0, 1).text = "Construccions del Vallès SL"
    table.cell(1, 0).text = "Vehicle"
    table.cell(1, 1).text = "Matrícula 1234 BCD, conductora Anna Vila"

    commented = doc.add_paragraph("Paràgraf amb comentari.")
    doc.add_comment(
        commented.runs,
        text="Confirmar amb Joan Puig i Ferrer",
        author="Montserrat Soler i Vidal",
        initials="MS",
    )

    doc.add_paragraph("La part és major d’edat i viu al carrer major del poble.")
    doc.save(str(path))
    return path


def build_demanda_es(path: Path) -> Path:
    doc = Document()
    doc.core_properties.author = "Laura Martín Ruiz"
    doc.core_properties.last_modified_by = "Laura Martín Ruiz"
    doc.core_properties.title = "Demanda Juan Pérez García contra Construcciones Levante SL"

    section = doc.sections[0]
    section.header.paragraphs[0].text = "Despacho Martín & Asociados — Letrada Laura Martín Ruiz"
    section.footer.paragraphs[0].text = "Tel. 612 345 678 · info@martin-asociados.example"

    doc.add_heading("AL JUZGADO DE PRIMERA INSTANCIA", level=1)

    p = doc.add_paragraph("Don ")
    p.add_run("Juan ").bold = True
    p.add_run("Pé").bold = True
    p.add_run("rez García").italic = True
    p.add_run(", mayor de edad, con DNI 12345678Z y domicilio en Calle Mayor 5, Valencia, ")
    p.add_run("actuando en su propio nombre, ante el Juzgado comparece y DICE:")

    doc.add_paragraph(
        "Que el Sr. Pérez suscribió contrato con CONSTRUCCIONES LEVANTE SL (CIF B12345674) "
        "en el procedimiento 1234/2025. J. Pérez García abonó las cantidades en la cuenta "
        "ES91 2100 0418 4502 0005 1332."
    )

    p = doc.add_paragraph("Contacto del demandante:\t")
    p.add_run("Juan")
    p.add_run().add_tab()
    p.add_run("Pérez García")

    p = doc.add_paragraph("Correo: ")
    _add_hyperlink(p, "mailto:juan.perez@correo.example", "juan.perez@correo.example")

    p = doc.add_paragraph("Texto revisado. ")
    _add_tracked_deletion(p, "Antes ponía Juan Pérez García", "Laura Martín Ruiz")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Demandado"
    table.cell(0, 1).text = "Construcciones Levante SL"
    table.cell(1, 0).text = "Vehículo"
    table.cell(1, 1).text = "Matrícula 1234 BCD, conductor Ana López"

    commented = doc.add_paragraph("Párrafo con comentario.")
    doc.add_comment(
        commented.runs,
        text="Confirmar con Juan Pérez García",
        author="Laura Martín Ruiz",
        initials="LM",
    )

    doc.add_paragraph("La parte es mayor de edad y vive en la calle mayor del pueblo.")
    doc.save(str(path))
    return path


PDF_LINES = [
    "AL JUTJAT DE PRIMERA INSTÀNCIA",
    "",
    "En Joan Puig i Ferrer, major d'edat, amb DNI 12345678Z i domicili al Carrer Major 5,",
    "Girona. Que el Sr. Puig va subscriure un contracte amb CONSTRUCCIONS DEL VALLÈS SL",
    "(NIF B12345674) en el procediment 1234/2025. El pagament el va fer J. Puig",
    "Ferrer al compte ES91 2100 0418 4502 0005 1332.",
    "La part contrària, Martí d'Orriols, no ha pagat. El perit Marcel·lí Roca ho confirma.",
    "La part és major d'edat i viu al carrer major del poble.",
]


SCANNED_LINES = [
    "ANNEX 1 - REBUT",
    "Jo, Joan Puig i Ferrer, amb DNI 12345678Z, he rebut de Martí d'Orriols",
    "la quantitat de 1.500 euros en concepte de bestreta.",
    "Girona, 3 de març de 2025.",
    "Signat: J. Puig Ferrer",
]


def build_demanda_pdf_ca(path: Path) -> Path:
    """PDF en catalán: página con texto + página escaneada (solo imagen)."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in PDF_LINES:
        page.insert_text((56, y), line, fontsize=10, fontname="helv")
        y += 18

    # Tabla con bordes (las líneas deben conservarse tras anonimizar)
    top = y + 10
    for i in range(3):
        page.draw_line((56, top + i * 22), (456, top + i * 22))
    for x in (56, 196, 456):
        page.draw_line((x, top), (x, top + 44))
    page.insert_text((62, top + 15), "Demandat", fontsize=10)
    page.insert_text((202, top + 15), "Construccions del Vallès SL", fontsize=10)
    page.insert_text((62, top + 37), "Vehicle", fontsize=10)
    page.insert_text((202, top + 37), "Matrícula 1234 BCD, conductora Anna Vila", fontsize=10)

    y = top + 70
    page.insert_text((56, y), "Correu: joan.puig@correu.example", fontsize=10)
    page.insert_link({
        "kind": pymupdf.LINK_URI, "from": pymupdf.Rect(96, y - 10, 240, y + 2),
        "uri": "mailto:joan.puig@correu.example",
    })  # fmt: skip
    page.insert_text(
        (56, 800), "Despatx Soler & Associats · Tel. 972 123 456 · info@soler-associats.example",
        fontsize=8,
    )  # fmt: skip
    note = page.add_text_annot((480, 72), "Confirmar amb Joan Puig i Ferrer")
    note.set_info(title="Montserrat Soler i Vidal")
    note.update()

    # Página "escaneada": una imagen con texto, sin capa de texto.
    scan_src = pymupdf.open()
    scan_page = scan_src.new_page()
    for i, line in enumerate(SCANNED_LINES):
        scan_page.insert_text((56, 90 + i * 22), line, fontsize=12, fontname="tiro")
    pix = scan_page.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
    scanned = doc.new_page()
    scanned.insert_image(scanned.rect, pixmap=pix)

    doc.set_metadata({"author": "Montserrat Soler i Vidal", "title": "Demanda Joan Puig i Ferrer"})
    doc.set_toc([[1, "Demanda de Joan Puig i Ferrer", 1], [1, "Annex", 2]])
    doc.save(str(path))
    return path


# --- Demanda andorrana -----------------------------------------------------------------
# Datos andorranos (pasaporte, NRT, CASS, residencia, +376, AD500...) junto a organismos
# que NO deben anonimizarse (Batllia, Comú, CASS como institución, BOPA, Cos de Policia).
ANDORRA_PARAGRAPHS = [
    "A LA BATLLIA D'ANDORRA",
    "En Pere Riba i Montané, major d'edat, de nacionalitat andorrana, amb passaport "
    "andorrà O1234567, afiliat a la CASS amb el número 123456 i domicili a l'Avinguda "
    "Meritxell, 25, 3r 2a, AD500 Andorra la Vella, telèfon +376 812 345, compareix i DIU:",
    "Primer. Que la Sra. Laura Fontanals Pujol, de nacionalitat portuguesa, amb passaport "
    "CB1234567 i permís de residència i treball núm. 12345/2020, és administradora "
    "d'ESQUÍS PIRINEU SL, amb NRT L-712345-B.",
    "Segon. Que el Sr. Riba va pagar les factures al compte AD12 0001 2030 2003 5910 0100 "
    "i va rebre el vehicle matrícula A 1234, que conduïa Núria Serra (mòbil 361 234).",
    "Tercer. Que la Sra. Fontanals té NRT F-654321-Z i ha estat citada per la causa 1234/2025. "
    "Correu del demandant: pere.riba@correu.example.",
    "Quart. Que la Caixa Andorrana de Seguretat Social va informar el 2024 que la CASS no "
    "té constància del pagament, segons el Butlletí Oficial del Principat d'Andorra. "
    "El Comú d'Andorra la Vella i el Cos de Policia ho poden confirmar.",
]
ANDORRA_HEADER = "Despatx Casals Advocats — Lletrat Jordi Casals i Pons"
ANDORRA_FOOTER = "Tel. 00376 712 345 · info@casals-advocats.example"


def build_demanda_andorra(path: Path) -> Path:
    """Demanda andorrana en catalán (Word), con datos inventados."""
    doc = Document()
    doc.core_properties.author = "Jordi Casals i Pons"
    doc.core_properties.title = "Demanda Pere Riba i Montané contra Esquís Pirineu SL"
    section = doc.sections[0]
    section.header.paragraphs[0].text = ANDORRA_HEADER
    section.footer.paragraphs[0].text = ANDORRA_FOOTER

    doc.add_heading(ANDORRA_PARAGRAPHS[0], level=1)
    for text in ANDORRA_PARAGRAPHS[1:]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Part demandada"
    table.cell(0, 1).text = "Esquís Pirineu SL (NRT L-712345-B)"
    table.cell(1, 0).text = "Telèfon"
    table.cell(1, 1).text = "00376 361 234"
    doc.save(str(path))
    return path


def build_demanda_andorra_pdf(path: Path) -> Path:
    """La misma demanda andorrana en PDF con texto."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((56, 40), ANDORRA_HEADER, fontsize=8)
    box = pymupdf.Rect(56, 72, 540, 760)
    html = "".join(f"<p>{t}</p>" for t in ANDORRA_PARAGRAPHS)
    html += "<p>Part demandada: Esquís Pirineu SL (NRT L-712345-B). Telèfon 00376 361 234</p>"
    page.insert_htmlbox(box, html, css="* {font-family: sans-serif; font-size: 10pt;}")
    page.insert_text((56, 800), ANDORRA_FOOTER, fontsize=8)
    doc.set_metadata({"author": "Jordi Casals i Pons", "title": "Demanda Pere Riba i Montané"})
    doc.save(str(path))
    return path


if __name__ == "__main__":
    here = Path(__file__).parent
    for out in (
        build_demanda_ca(here / "demanda_ficticia_ca.docx"),
        build_demanda_pdf_ca(here / "demanda_ficticia_ca.pdf"),
        build_demanda_es(here / "demanda_ficticia_es.docx"),
        build_demanda_andorra(here / "demanda_andorra.docx"),
        build_demanda_andorra_pdf(here / "demanda_andorra.pdf"),
    ):
        print(f"Creado {out}")
