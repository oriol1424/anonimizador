"""Regresiones de la auditoría del 2026-10-11 (docs/AUDITORIA.md). Datos inventados."""

import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches

from anonimizador.core.matcher import Matcher, Term
from anonimizador.core.pipeline import anonymize
from anonimizador.core.regex_es import detect

NAME = "Joan Puig i Ferrer"
TERMS = [Term(NAME, "[CLIENT_1]")]


def _zip_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())


def _word_with_extras(path: Path) -> Path:
    """Word con texto alternativo, propiedad personalizada y miniatura (como los de Word)."""
    base = path.with_name("base.docx")
    doc = Document()
    doc.add_paragraph(f"En {NAME} compareix.")
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.clear_with(200)
    image = path.with_name("foto.png")
    pix.save(str(image))
    doc.add_picture(str(image), width=Inches(1))
    doc.inline_shapes[0]._inline.docPr.set("descr", f"Foto del DNI de {NAME}")
    doc.save(str(base))
    custom = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="http://'
        'schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://'
        'schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CD'
        'D505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Client">'
        f"<vt:lpwstr>{NAME}</vt:lpwstr></property></Properties>"
    )
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/docProps/custom.xml" '
                                    b'ContentType="application/vnd.openxmlformats-officedocument'
                                    b'.custom-properties+xml"/></Types>')  # fmt: skip
            if item.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdC" Type="http:'
                                    b'//schemas.openxmlformats.org/officeDocument/2006/relation'
                                    b'ships/custom-properties" Target="docProps/custom.xml"/>'
                                    b"</Relationships>")  # fmt: skip
            zout.writestr(item, data)
        zout.writestr("docProps/custom.xml", custom)
    return path


@pytest.fixture
def word_out(tmp_path):
    src = _word_with_extras(tmp_path / "doc.docx")
    assert "thumbnail" in _zip_text(src) and "Foto del DNI de Joan" in _zip_text(src)
    out = tmp_path / "out.docx"
    return out, anonymize(src, out, TERMS)


def test_word_alt_text_custom_props_and_thumbnail(word_out):
    out, report = word_out
    text = _zip_text(out)
    assert "Puig" not in text  # ni texto alternativo ni propiedad personalizada
    assert "Foto del DNI de [CLIENT_1]" in text
    with zipfile.ZipFile(out) as z:
        assert not [n for n in z.namelist() if "thumbnail" in n]
    assert not [n for n in report.notices if n.is_leak]


def test_pdf_text_outside_cropbox_is_removed(tmp_path):
    src = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((56, 72), f"Fora del recorte: {NAME}")
    page.insert_text((56, 400), f"Visible: {NAME}")
    page.set_cropbox(pymupdf.Rect(0, 200, 595, 842))
    doc.save(str(src))
    out = tmp_path / "out.pdf"
    anonymize(src, out, TERMS)
    result = pymupdf.open(str(out))[0]
    result.set_cropbox(result.mediabox)
    assert "Puig" not in result.get_text()


def test_pdf_attachments_are_removed(tmp_path):
    src = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((56, 72), f"En {NAME} compareix.")
    page.add_file_annot((300, 72), f"DNI de {NAME}".encode(), "nota.txt")
    doc.embfile_add("dades.txt", f"DNI de {NAME}: 12345678Z".encode())
    doc.save(str(src))
    out = tmp_path / "out.pdf"
    report = anonymize(src, out, TERMS)
    result = pymupdf.open(str(out))
    assert result.embfile_count() == 0
    assert not list(result[0].annots())
    assert [n.count for n in report.notices if n.kind == "attachments_removed"] == [2]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("el Sr. Mas va dir", "el Sr. [CLIENT_1] va dir"),
        ("la Sra. Mas i Pons", "la Sra. [CLIENT_1]"),
        ("En Mas va signar", "En [CLIENT_1] va signar"),
        ("un mas antic", "un mas antic"),  # palabra corriente
        ("Mas de Can Pons", "Mas de Can [CLIENT_1]"),  # sin tratamiento: no
    ],
)
def test_short_surnames_after_a_title(text, expected):
    assert Matcher([Term("Pere Mas i Pons", "[CLIENT_1]")]).replace(text) == expected


@pytest.mark.parametrize(
    ("text", "label", "value"),
    [
        ("DNI 12345678z", "DNI", "12345678z"),
        ("NIE x1234567l", "NIE", "x1234567l"),
        ("tel. 81 23 45", "TELEFONO", "81 23 45"),
        ("tel. +376 81 23 45", "TELEFONO", "+376 81 23 45"),
        ("tel. (+376) 812345", "TELEFONO", "(+376) 812345"),
        ("mòbil 376 612 345", "TELEFONO", "376 612 345"),
        ("Data de naixement: 03/04/1980", "NACIMIENTO", "03/04/1980"),
        ("nascut el 3 de març de 1980", "NACIMIENTO", "3 de març de 1980"),
        ("nacida el 12 de mayo de 1975", "NACIMIENTO", "12 de mayo de 1975"),
        ("causa 0001234/2025", "PROCEDIMIENTO", "0001234/2025"),
    ],
)
def test_formats_found_in_the_audit(text, label, value):
    assert (label, value) in [(d.label, d.text) for d in detect(text)]


@pytest.mark.parametrize(
    "text", ["pagar 12345678 a la part", "el 03/04/1980 es va signar", "import de 600.000 euros"]
)
def test_no_false_positives_from_new_rules(text):
    assert [d.label for d in detect(text)] == []
