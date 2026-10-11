"""Imágenes dentro de Word y PDF (p. ej. un DNI escaneado pegado): OCR y borrado."""

import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches

from anonimizador.core import images
from anonimizador.core.lang import CATALAN
from anonimizador.core.matcher import Term
from anonimizador.core.pipeline import analyze, anonymize

NAME = "Joan Puig i Ferrer"
TERMS = [Term(NAME, "[CLIENT_1]"), Term("12345678Z", "[DNI_1]", expand=False)]


def _photo(path: Path) -> Path:
    """«Foto» de un documento: texto dibujado en una imagen PNG, sin capa de texto."""
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=120)
    page.insert_text((20, 45), f"Titular: {NAME}", fontsize=18, fontname="tiro")
    page.insert_text((20, 85), "DNI 12345678Z", fontsize=18, fontname="tiro")
    page.get_pixmap(dpi=150, alpha=False).save(str(path))
    return path


def _media_text(docx: Path) -> str:
    with zipfile.ZipFile(docx) as z:
        return " ".join(images.text_of(z.read(n)) for n in z.namelist() if "/media/" in n)


@pytest.fixture
def word_with_photo(tmp_path: Path) -> Path:
    src = tmp_path / "amb_foto.docx"
    doc = Document()
    doc.add_paragraph("Annex: còpia del document d'identitat.")
    doc.add_picture(str(_photo(tmp_path / "foto.png")), width=Inches(4))
    doc.save(str(src))
    return src


def test_names_in_word_images_are_suggested(word_with_photo):
    forms = {f for s in analyze(word_with_photo, [], CATALAN).suggestions for f in s.forms}
    assert "12345678Z" in forms


def test_word_image_is_anonymized(word_with_photo, tmp_path):
    assert "Puig" in _media_text(word_with_photo)
    out = tmp_path / "out.docx"
    report = anonymize(word_with_photo, out, TERMS)
    text = _media_text(out)
    assert "Puig" not in text and "12345678" not in text
    assert "CLIENT" in text  # el alias está escrito en la imagen
    assert [n.count for n in report.notices if n.kind == "images_ocr"] == [1]
    assert not [n for n in report.notices if n.is_leak]


def test_pdf_image_on_a_text_page_is_anonymized(tmp_path):
    src = tmp_path / "amb_foto.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((56, 72), "Annex: còpia del document d'identitat.")
    page.insert_image(pymupdf.Rect(56, 100, 476, 220), filename=str(_photo(tmp_path / "f.png")))
    doc.save(str(src))
    forms = {f for s in analyze(src, [], CATALAN).suggestions for f in s.forms}
    assert "12345678Z" in forms

    out = tmp_path / "out.pdf"
    report = anonymize(src, out, TERMS)
    result = pymupdf.open(str(out))[0]
    seen = images.read(result.get_pixmap(dpi=200, clip=pymupdf.Rect(56, 100, 476, 220)))[0]
    assert "Puig" not in seen and "12345678" not in seen
    assert "Annex" in result.get_text()  # el texto de la página sigue
    assert not [n for n in report.notices if n.is_leak]


def test_small_icons_are_ignored():
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 30, 30), False)
    assert not images.worth_reading(pix)


def test_unreadable_formats_are_reported(tmp_path):
    assert images.load(b"\x01\x00\x00\x00 no es una imatge") is None
